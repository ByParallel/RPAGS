"""
FIAP Global Solution 2026 - AI for RPA
Sistema Automatizado de Analise de Missoes Espaciais

Pipeline:
  1. NLP: processa logs de missao -> extrai eventos, severidade, keywords
  2. Visao Computacional: analisa imagens de sensores -> detecta anomalias
  3. Relatorio: gera Excel consolidado com todos os resultados
"""
import os
import glob
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from nlp_module import analisar_log
from vision_module import analisar_imagem

# ── Configuracoes ────────────────────────────────────────────────────────────
DIR_LOGS    = "data/logs"
DIR_IMAGENS = "data/images"
DIR_OUTPUT  = "outputs"
ARQUIVO_SAIDA = os.path.join(DIR_OUTPUT, "relatorio_missoes.xlsx")

# Cores para o Excel
COR_CRITICO = "FF4444"
COR_ALERTA  = "FFA500"
COR_NORMAL  = "44BB44"
COR_INFO    = "4488FF"
COR_HEADER  = "1A237E"
COR_TITULO  = "0D47A1"


# ── Utilitarios ──────────────────────────────────────────────────────────────
def cor_por_status(status: str) -> str:
    mapa = {"CRITICO": COR_CRITICO, "ALERTA": COR_ALERTA,
            "NORMAL": COR_NORMAL, "INFO": COR_INFO}
    return mapa.get(status.upper(), "FFFFFF")


def aplicar_header(ws, linha: int, valores: list, cor_fundo: str = COR_HEADER):
    fill = PatternFill("solid", fgColor=cor_fundo)
    font = Font(color="FFFFFF", bold=True)
    for col, valor in enumerate(valores, 1):
        cell = ws.cell(row=linha, column=col, value=valor)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def autoajustar_colunas(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


# ── Etapa 1: Processar logs com NLP ─────────────────────────────────────────
def processar_logs() -> list[dict]:
    arquivos = glob.glob(os.path.join(DIR_LOGS, "*.txt"))
    if not arquivos:
        print("  [AVISO] Nenhum arquivo de log encontrado em", DIR_LOGS)
        return []

    resultados = []
    for arquivo in sorted(arquivos):
        nome = os.path.basename(arquivo)
        print(f"  Analisando log: {nome}")
        resultado = analisar_log(arquivo)
        resultado["arquivo"] = nome
        resultados.append(resultado)
        print(f"    Missao: {resultado['missao']} | Status: {resultado['status_geral']} "
              f"| Criticos: {resultado['criticos']} | Alertas: {resultado['alertas']}")

    return resultados


# ── Etapa 2: Processar imagens com Visao Computacional ──────────────────────
def processar_imagens() -> list[dict]:
    extensoes = ["*.png", "*.jpg", "*.jpeg"]
    arquivos = []
    for ext in extensoes:
        arquivos.extend(glob.glob(os.path.join(DIR_IMAGENS, ext)))

    if not arquivos:
        print("  [AVISO] Nenhuma imagem encontrada em", DIR_IMAGENS)
        print("  Execute: python gerar_imagens.py")
        return []

    resultados = []
    for arquivo in sorted(arquivos):
        nome = os.path.basename(arquivo)
        print(f"  Analisando imagem: {nome}")
        try:
            resultado = analisar_imagem(arquivo)
            resultados.append(resultado)
            anomalia = "SIM" if resultado["anomalia_detectada"] else "NAO"
            print(f"    Classificacao: {resultado['classificacao']} | Anomalia: {anomalia}")
        except Exception as e:
            print(f"    [ERRO] {e}")

    return resultados


# ── Etapa 3: Gerar relatorio Excel ──────────────────────────────────────────
def gerar_relatorio(dados_nlp: list[dict], dados_visao: list[dict]):
    os.makedirs(DIR_OUTPUT, exist_ok=True)
    wb = openpyxl.Workbook()

    # ── Aba 1: Resumo Executivo ──────────────────────────────────────────────
    ws_resumo = wb.active
    ws_resumo.title = "Resumo Executivo"

    total_missoes   = len(dados_nlp)
    total_criticos  = sum(d["criticos"] for d in dados_nlp)
    total_alertas   = sum(d["alertas"] for d in dados_nlp)
    total_imagens   = len(dados_visao)
    imagens_anomalas = sum(1 for d in dados_visao if d["anomalia_detectada"])

    ws_resumo.merge_cells("A1:F1")
    titulo = ws_resumo["A1"]
    titulo.value = "SISTEMA DE ANALISE DE MISSOES ESPACIAIS - FIAP GS 2026"
    titulo.font = Font(color="FFFFFF", bold=True, size=14)
    titulo.fill = PatternFill("solid", fgColor=COR_TITULO)
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    ws_resumo.row_dimensions[1].height = 30

    ws_resumo["A3"] = "Data de Geracao:"
    ws_resumo["B3"] = datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ws_resumo["A4"] = "Total de Missoes Analisadas:"
    ws_resumo["B4"] = total_missoes
    ws_resumo["A5"] = "Total de Eventos Criticos:"
    ws_resumo["B5"] = total_criticos
    ws_resumo["B5"].font = Font(color=COR_CRITICO, bold=True)
    ws_resumo["A6"] = "Total de Alertas:"
    ws_resumo["B6"] = total_alertas
    ws_resumo["B6"].font = Font(color=COR_ALERTA, bold=True)
    ws_resumo["A7"] = "Total de Imagens Analisadas:"
    ws_resumo["B7"] = total_imagens
    ws_resumo["A8"] = "Imagens com Anomalia Detectada:"
    ws_resumo["B8"] = imagens_anomalas

    for row in range(3, 9):
        ws_resumo.cell(row=row, column=1).font = Font(bold=True)

    # Tabela de missoes no resumo
    aplicar_header(ws_resumo, 10,
                   ["Missao", "Data", "Status Geral", "Criticos", "Alertas", "Palavras-chave"])

    for i, d in enumerate(dados_nlp, 11):
        ws_resumo.cell(row=i, column=1, value=d["missao"])
        ws_resumo.cell(row=i, column=2, value=d["data"])
        status_cell = ws_resumo.cell(row=i, column=3, value=d["status_geral"])
        status_cell.fill = PatternFill("solid", fgColor=cor_por_status(d["status_geral"]))
        status_cell.font = Font(color="FFFFFF", bold=True)
        ws_resumo.cell(row=i, column=4, value=d["criticos"])
        ws_resumo.cell(row=i, column=5, value=d["alertas"])
        ws_resumo.cell(row=i, column=6, value=", ".join(d["palavras_chave"][:5]))

    autoajustar_colunas(ws_resumo)

    # ── Aba 2: Analise NLP Detalhada ────────────────────────────────────────
    ws_nlp = wb.create_sheet("Analise NLP - Logs")
    aplicar_header(ws_nlp, 1,
                   ["Missao", "Hora", "Severidade", "Descricao do Evento"])

    linha = 2
    for d in dados_nlp:
        for evento in d["eventos"]:
            ws_nlp.cell(row=linha, column=1, value=d["missao"])
            ws_nlp.cell(row=linha, column=2, value=evento["hora"])
            sev_cell = ws_nlp.cell(row=linha, column=3, value=evento["severidade"])
            sev_cell.fill = PatternFill("solid", fgColor=cor_por_status(evento["severidade"]))
            sev_cell.font = Font(color="FFFFFF", bold=True)
            ws_nlp.cell(row=linha, column=4, value=evento["descricao"])
            linha += 1

    autoajustar_colunas(ws_nlp)

    # ── Aba 3: Palavras-chave por Missao ────────────────────────────────────
    ws_kw = wb.create_sheet("Palavras-chave")
    aplicar_header(ws_kw, 1, ["Missao", "Status", "Top 10 Palavras-chave"])

    for i, d in enumerate(dados_nlp, 2):
        ws_kw.cell(row=i, column=1, value=d["missao"])
        st = ws_kw.cell(row=i, column=2, value=d["status_geral"])
        st.fill = PatternFill("solid", fgColor=cor_por_status(d["status_geral"]))
        st.font = Font(color="FFFFFF", bold=True)
        ws_kw.cell(row=i, column=3, value=" | ".join(d["palavras_chave"]))

    autoajustar_colunas(ws_kw)

    # ── Aba 4: Visao Computacional ──────────────────────────────────────────
    ws_vis = wb.create_sheet("Visao Computacional")
    headers_vis = [
        "Arquivo", "Resolucao", "Brilho Medio", "Desvio Padrao", "Contraste",
        "Regioes Anomalas", "% Area Anomala", "Anomalia?", "Classificacao",
        "% Pixels Escuros", "% Pixels Medios", "% Pixels Brilhantes"
    ]
    aplicar_header(ws_vis, 1, headers_vis)

    for i, d in enumerate(dados_visao, 2):
        ws_vis.cell(row=i, column=1,  value=d["arquivo"])
        ws_vis.cell(row=i, column=2,  value=d["resolucao"])
        ws_vis.cell(row=i, column=3,  value=d["brilho_medio"])
        ws_vis.cell(row=i, column=4,  value=d["desvio_padrao"])
        ws_vis.cell(row=i, column=5,  value=d["contraste"])
        ws_vis.cell(row=i, column=6,  value=d["regioes_anomalas"])
        ws_vis.cell(row=i, column=7,  value=d["percentual_anomalo"])
        anomalia_val = "SIM" if d["anomalia_detectada"] else "NAO"
        an_cell = ws_vis.cell(row=i, column=8, value=anomalia_val)
        an_cell.fill = PatternFill("solid",
                                   fgColor=COR_CRITICO if d["anomalia_detectada"] else COR_NORMAL)
        an_cell.font = Font(color="FFFFFF", bold=True)
        ws_vis.cell(row=i, column=9,  value=d["classificacao"])
        ws_vis.cell(row=i, column=10, value=d["pixels_escuros_pct"])
        ws_vis.cell(row=i, column=11, value=d["pixels_medios_pct"])
        ws_vis.cell(row=i, column=12, value=d["pixels_brilhantes_pct"])

    autoajustar_colunas(ws_vis)

    # ── Aba 5: Eventos Criticos (resumo rapido) ──────────────────────────────
    ws_crit = wb.create_sheet("Eventos Criticos")
    aplicar_header(ws_crit, 1, ["Missao", "Data", "Descricao do Evento Critico"])

    linha = 2
    for d in dados_nlp:
        for desc in d["detalhes_criticos"]:
            ws_crit.cell(row=linha, column=1, value=d["missao"])
            ws_crit.cell(row=linha, column=2, value=d["data"])
            cell = ws_crit.cell(row=linha, column=3, value=desc)
            cell.fill = PatternFill("solid", fgColor="FFE0E0")
            linha += 1

    if linha == 2:
        ws_crit.cell(row=2, column=1, value="Nenhum evento critico registrado.")

    autoajustar_colunas(ws_crit)

    wb.save(ARQUIVO_SAIDA)
    print(f"\n  Relatorio salvo em: {ARQUIVO_SAIDA}")


# ── Pipeline principal ───────────────────────────────────────────────────────
def main():
    print("=" * 60)
    print(" FIAP GS 2026 - Sistema de Analise de Missoes Espaciais")
    print("=" * 60)

    print("\n[1/3] Processando logs de missao com NLP...")
    dados_nlp = processar_logs()

    print("\n[2/3] Analisando imagens com Visao Computacional...")
    dados_visao = processar_imagens()

    if not dados_nlp and not dados_visao:
        print("\n[ERRO] Sem dados para processar. Verifique as pastas data/logs e data/images.")
        return

    print("\n[3/3] Gerando relatorio Excel...")
    gerar_relatorio(dados_nlp, dados_visao)

    print("\n" + "=" * 60)
    print(" PIPELINE CONCLUIDO COM SUCESSO!")
    print(f" Logs analisados : {len(dados_nlp)}")
    print(f" Imagens analisadas: {len(dados_visao)}")
    print(f" Relatorio: {ARQUIVO_SAIDA}")
    print("=" * 60)


if __name__ == "__main__":
    main()
